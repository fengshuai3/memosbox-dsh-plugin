"""One bounded local-library request. No server, model clients, .env or parse_file.

The parent applies OS confinement; the Python audit hook is defense in depth,
not an OS sandbox. Failures never print source text, paths or tracebacks.
"""
from __future__ import annotations

import errno
import hashlib
import io
import json
import logging
import os
from pathlib import Path
import platform
import re
import stat
import sys
import zipfile

SCHEMA_VERSION = 1
MIROBODY_VERSION = "1.4.2"
BUNDLE_VERSION = "loinc-2.82+2026.08.28-af2524b7a285"
OPERATIONS = ("status", "resolve", "normalize_readings", "metric_info", "extract_document")
MAX_REQUEST = 2 * 1024 * 1024
MAX_TEXT = 100_000


class WorkerError(Exception):
    pass


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def bounded_string(value, *, limit=4096, optional=False) -> str:
    if value is None and optional:
        return ""
    if not isinstance(value, str) or len(value) > limit or "\x00" in value:
        raise WorkerError("INVALID_STRING")
    return value


def limit_value(limits, name, default, maximum):
    value = limits.get(name, default)
    if type(value) is not int or not 1 <= value <= maximum:
        raise WorkerError("INVALID_LIMIT")
    return value


def prohibit_external_effects(event, args):
    if event.startswith("socket.") or event in {"subprocess.Popen", "os.system", "os.exec", "os.posix_spawn", "os.fork", "os.forkpty"}:
        raise PermissionError(errno.EPERM, "WORKER_EXTERNAL_EFFECT_DENIED")


def probe(payload):
    """Runs BEFORE the Python hook: EPERM/EACCES must originate in the OS."""
    import socket
    allowed = Path(payload["allowedPath"]).read_text() == "allowed-synthetic-probe"
    denied = False
    try:
        Path(payload["deniedPath"]).read_bytes()
    except OSError as exc:
        denied = exc.errno in (errno.EPERM, errno.EACCES)
    results = []
    for family, socktype, destination in (
        (socket.AF_INET, socket.SOCK_STREAM, ("127.0.0.1", 9)),
        (socket.AF_INET, socket.SOCK_DGRAM, ("192.0.2.1", 9)),
        (socket.AF_INET6, socket.SOCK_STREAM, ("::1", 9)),
    ):
        sock = None
        try:
            sock = socket.socket(family, socktype)
            sock.settimeout(1)
            sock.connect(destination)
            results.append(False)
        except OSError as exc:
            results.append(exc.errno in (errno.EPERM, errno.EACCES))
        finally:
            if sock is not None:
                sock.close()
    return {"schemaVersion": 1, "probe": {"allowedRead": allowed, "deniedRead": denied, "networkDenied": all(results), "networkChecks": results, "environmentNames": sorted(os.environ)}}


def versions():
    import mirobody
    from mirobody.kernel import metrics
    if mirobody.__version__ != MIROBODY_VERSION or mirobody.BUNDLE_VERSION != BUNDLE_VERSION:
        raise WorkerError("RUNTIME_VERSION_MISMATCH")
    root = Path(mirobody.__file__).parent
    return {"pythonVersion": platform.python_version(), "mirobodyVersion": mirobody.__version__,
            "bundleSha256": sha((root / "res" / "fhir_loinc_bundle.tar.gz").read_bytes()),
            "deviceCatalogVersion": metrics.TERMINOLOGY_VERSION,
            "deviceCatalogSha256": sha((root / "res" / "metrics.tsv").read_bytes()),
            "network": "blocked", "isolation": "unverified", "supportedOperations": list(OPERATIONS)}


def normalize_reading(raw):
    from mirobody import resolve_reading
    from mirobody.units import normalize_unit, parse_value_unit, unit_families
    if not isinstance(raw, dict):
        raise WorkerError("INVALID_READING")
    name = bounded_string(raw.get("rawName", raw.get("name")), limit=512)
    value = bounded_string(raw.get("rawValue", raw.get("value")), optional=True)
    unit = bounded_string(raw.get("rawUnit", raw.get("unit")), limit=256, optional=True)
    refs = raw.get("evidenceRefs", [])
    if not isinstance(refs, list) or len(refs) > 100:
        raise WorkerError("INVALID_EVIDENCE")
    refs = [bounded_string(ref, limit=160) for ref in refs]
    warnings = []
    if not name.strip():
        raise WorkerError("EMPTY_READING_NAME")
    if re.fullmatch(r"[+-]?(?:inf(?:inity)?|nan)", value.strip(), flags=re.I):
        warnings.append("NON_FINITE_VALUE_PRESERVED_NOT_CONVERTED")
    # Never serialize ParsedQuantity.value: it is a float and loses operators,
    # scientific spelling, ranges, and precision. Only reuse the unit parser.
    parsed = parse_value_unit(value) if value else None
    inferred_unit = parsed.unit if parsed else ""
    normalized_unit = normalize_unit(unit or inferred_unit or "") or ""
    if (unit or inferred_unit) and (not normalized_unit or not unit_families(normalized_unit)):
        warnings.append("UNIT_UNRECOGNIZED")
    if unit and inferred_unit and normalize_unit(inferred_unit) != normalize_unit(unit):
        warnings.append("INLINE_UNIT_CONFLICT")
    hit = resolve_reading(name, value or None, unit or None)
    state = "refused" if hit.method == "refused" else "ready" if hit.resolved and hit.loinc else "unresolved"
    unsafe_reading = any(code in warnings for code in ("NON_FINITE_VALUE_PRESERVED_NOT_CONVERTED", "INLINE_UNIT_CONFLICT", "UNIT_UNRECOGNIZED"))
    if unsafe_reading and state != "refused":
        state = "unresolved"
    if hit.candidates > 1:
        warnings.append("MULTIPLE_LEXICAL_CANDIDATES_REVIEW_REQUIRED")
    return {"rawName": name, "rawValue": value, "rawUnit": unit, "normalizedValue": value,
            "normalizedUnit": normalized_unit, "status": state, "resolutionMethod": hit.method,
            "codeSystem": "loinc" if hit.loinc and not unsafe_reading else "", "code": hit.loinc if not unsafe_reading else "", "canonical": hit.canonical if not unsafe_reading else "",
            "candidates": hit.candidates, "bundleVersion": BUNDLE_VERSION, "evidenceRefs": refs,
            "warnings": warnings}


def segment(text, location, state="ready"):
    digest = sha(text.encode("utf-8"))
    identity = sha(json.dumps(location, sort_keys=True).encode() + bytes.fromhex(digest))[:24]
    return {"id": "seg-" + identity, "text": text, "location": {**location, "bbox": "unavailable"}, "status": state, "digest": digest}


def extract_document(payload, limits):
    from mirobody.documents import extract
    source = bounded_string(payload.get("sourcePath"), limit=4096)
    expected = bounded_string(payload.get("sourceDigest"), limit=64)
    if not re.fullmatch(r"[a-f0-9]{64}", expected):
        raise WorkerError("SOURCE_DIGEST_REQUIRED")
    max_bytes = limit_value(limits, "maxFileBytes", 20 * 1024 * 1024, 20 * 1024 * 1024)
    fd = os.open(source, os.O_RDONLY | os.O_NOFOLLOW)
    try:
        facts = os.fstat(fd)
        if not stat.S_ISREG(facts.st_mode) or facts.st_size > max_bytes:
            raise WorkerError("SOURCE_TYPE_OR_SIZE")
        with os.fdopen(fd, "rb", closefd=False) as handle:
            data = handle.read(max_bytes + 1)
    finally:
        os.close(fd)
    if len(data) > max_bytes or sha(data) != expected:
        raise WorkerError("SOURCE_DIGEST_OR_SIZE_MISMATCH")
    kind = payload.get("format")
    segments = []
    warnings = []
    truncated = False
    char_budget = MAX_TEXT
    if kind == "text":
        if b"\x00" in data or data.startswith((b"%PDF-", b"PK\x03\x04")):
            raise WorkerError("FORMAT_SIGNATURE_MISMATCH")
        # Invoke the pinned public decoder without its stripping changing line
        # numbers: decode each original physical line independently.
        for number, raw_line in enumerate(data.splitlines(), 1):
            text = extract.decode_text(raw_line, cap=len(raw_line) + 1)
            if len(text) > char_budget:
                text = text[:char_budget]
                truncated = True
            if text:
                segments.append(segment(text, {"line": number}))
                char_budget -= len(text)
            if char_budget <= 0:
                truncated = True
                break
    elif kind == "pdf":
        if not data.startswith(b"%PDF-"):
            raise WorkerError("FORMAT_SIGNATURE_MISMATCH")
        import pypdfium2 as pdfium
        max_pages = limit_value(limits, "maxPdfPages", 100, 100)
        document = pdfium.PdfDocument(data)
        try:
            total = len(document)
            for index in range(min(total, max_pages)):
                page = document[index]
                text_page = None
                try:
                    text_page = page.get_textpage()
                    # Skip upstream rendering for empty pages (OCR is a
                    # separate DSH permission and has not occurred here).
                    if text_page.count_chars() == 0:
                        segments.append(segment("", {"page": index + 1}, "needs_ocr"))
                        continue
                    # Pinned PUBLIC API on one bounded page; do not split its
                    # output on arbitrary document text that resembles headers.
                    one = pdfium.PdfDocument.new()
                    try:
                        one.import_pages(document, [index])
                        buffer = io.BytesIO()
                        one.save(buffer)
                        extracted = extract.pdf_text_layer(buffer.getvalue())
                    finally:
                        one.close()
                    text = extracted.removeprefix("--- page 1 ---\n")
                    state = "ready" if len(text.strip()) >= extract.MIN_PAGE_TEXT else "needs_ocr"
                    if len(text) > char_budget:
                        text = text[:char_budget]
                        truncated = True
                    segments.append(segment(text, {"page": index + 1}, state))
                    char_budget -= len(text)
                except Exception:
                    segments.append(segment("", {"page": index + 1}, "failed"))
                    warnings.append("PDF_PAGE_EXTRACTION_FAILED")
                finally:
                    if text_page is not None:
                        text_page.close()
                    page.close()
                if char_budget <= 0:
                    truncated = True
                    break
            if total > max_pages:
                truncated = True
                warnings.append("PDF_PAGE_LIMIT")
        finally:
            document.close()
    elif kind == "xlsx":
        if not data.startswith(b"PK\x03\x04"):
            raise WorkerError("FORMAT_SIGNATURE_MISMATCH")
        # Bound expansion before openpyxl reads ZIP members. Never extract ZIP
        # names to disk. XML entities are blocked by defusedxml in openpyxl.
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2048 or sum(e.file_size for e in entries) > 100 * 1024 * 1024 or any(e.flag_bits & 1 for e in entries):
                raise WorkerError("XLSX_ARCHIVE_LIMIT")
            if "xl/workbook.xml" not in archive.namelist():
                raise WorkerError("FORMAT_SIGNATURE_MISMATCH")
        from openpyxl import load_workbook
        from openpyxl.xml import DEFUSEDXML
        if not DEFUSEDXML:
            raise WorkerError("XML_PROTECTION_UNAVAILABLE")
        max_rows = limit_value(limits, "maxRows", 5000, 5000)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        rows_seen = 0
        try:
            for sheet in workbook.worksheets:
                # Upstream xlsx_sheets drops blank rows and loses coordinates;
                # use the SAME underlying library with original cell positions.
                for row in sheet.iter_rows(max_col=256, max_row=max_rows + 1):
                    rows_seen += 1
                    if rows_seen > max_rows:
                        truncated = True
                        break
                    for cell in row:
                        if cell.value is None:
                            continue
                        if cell.data_type == "f":
                            warnings.append("XLSX_FORMULA_NOT_EVALUATED")
                        text = str(cell.value)
                        if len(text) > char_budget:
                            text = text[:char_budget]
                            truncated = True
                        entry = segment(text, {"sheet": sheet.title, "row": cell.row, "column": cell.column})
                        entry["provenance"] = {"cellDataType": cell.data_type, "numberFormat": cell.number_format, "originalDisplayAvailable": False}
                        if cell.data_type in ("n", "d"):
                            warnings.append("XLSX_NUMERIC_OR_DATE_VALUE_NOT_ORIGINAL_DISPLAY")
                        segments.append(entry)
                        char_budget -= len(text)
                        if char_budget <= 0:
                            truncated = True
                            break
                    if char_budget <= 0:
                        break
                if (sheet.max_column or 0) > 256:
                    warnings.append("XLSX_COLUMN_LIMIT")
                    truncated = True
                if rows_seen > max_rows or char_budget <= 0:
                    break
        finally:
            workbook.close()
    else:
        raise WorkerError("UNSUPPORTED_DOCUMENT_FORMAT")
    states = {entry["status"] for entry in segments}
    if not segments:
        warnings.append("NO_READABLE_CONTENT")
    if "needs_ocr" in states:
        warnings.append("OCR_NOT_PERFORMED")
    state = "partial" if truncated or "failed" in states or ("needs_ocr" in states and "ready" in states) else "needs_ocr" if "needs_ocr" in states else "ready" if segments else "unresolved"
    return {"status": state, "segments": segments, "sourceDigest": expected, "warnings": sorted(set(warnings)), "truncated": truncated}


def dispatch(request):
    if not isinstance(request, dict) or request.get("schemaVersion") != 1:
        raise WorkerError("INVALID_PROTOCOL")
    operation = request.get("operation")
    payload = request.get("payload", {})
    if operation == "_probe":
        return probe(payload)
    if operation not in OPERATIONS or not isinstance(payload, dict):
        raise WorkerError("UNSUPPORTED_OPERATION")
    sys.addaudithook(prohibit_external_effects)
    logging.disable(logging.CRITICAL)
    limits = request.get("limits", {})
    runtime = versions()
    result = {"schemaVersion": 1, "operationId": bounded_string(request.get("operationId"), limit=64),
              "operation": operation, "status": "ready", "bundleVersion": BUNDLE_VERSION, "warnings": [], "truncated": False}
    if operation == "status":
        result["runtime"] = runtime
    elif operation in ("resolve", "normalize_readings"):
        readings = [payload] if operation == "resolve" else payload.get("readings")
        if not isinstance(readings, list) or not 1 <= len(readings) <= limit_value(limits, "maxBatch", 256, 256):
            raise WorkerError("READING_BATCH_LIMIT")
        result["readings"] = [normalize_reading(raw) for raw in readings]
        states = {item["status"] for item in result["readings"]}
        result["status"] = next(iter(states)) if len(states) == 1 else "partial"
    elif operation == "metric_info":
        from dataclasses import asdict
        from mirobody.kernel import metrics
        name = bounded_string(payload.get("name", payload.get("rawName")), limit=512)
        metric = metrics.METRICS.get(name) or metrics.MEMBERS.get(name)
        result["metric"] = ({**asdict(metric), "canonical": list(metric.canonical), "terminologyVersion": metrics.TERMINOLOGY_VERSION} if metric else None)
        if not metric:
            result["status"] = "unresolved"
    else:
        result.update(extract_document(payload, limits))
    return result


def main():
    request = {}
    try:
        raw = sys.stdin.buffer.read(MAX_REQUEST + 1)
        if len(raw) > MAX_REQUEST:
            raise WorkerError("REQUEST_LIMIT")
        request = json.loads(raw)
        result = dispatch(request)
    except Exception as exc:
        code = str(exc) if isinstance(exc, WorkerError) else "LOCAL_EXTRACTION_FAILED"
        result = {"schemaVersion": 1, "operationId": request.get("operationId", "") if isinstance(request, dict) else "",
                  "operation": request.get("operation", "status") if isinstance(request, dict) else "status",
                  "status": "failed", "bundleVersion": BUNDLE_VERSION, "warnings": [], "truncated": False,
                  "error": {"code": code, "message": "Local worker failed; no source content was logged."}}
    sys.stdout.write(json.dumps(result, ensure_ascii=False, allow_nan=False, separators=(",", ":")) + "\n")
    return 1 if result.get("status") == "failed" else 0


if __name__ == "__main__":
    raise SystemExit(main())
